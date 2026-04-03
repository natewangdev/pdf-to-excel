from __future__ import annotations

from statistics import median
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment


def _group_close_values(values: List[float], tol: float) -> List[float]:
    if not values:
        return []
    sorted_vals = sorted(values)
    groups: List[List[float]] = [[sorted_vals[0]]]
    for v in sorted_vals[1:]:
        if abs(v - groups[-1][-1]) <= tol:
            groups[-1].append(v)
        else:
            groups.append([v])
    return [sum(g) / len(g) for g in groups]


def _nearest_index(values: List[float], target: float) -> int:
    return min(range(len(values)), key=lambda i: abs(values[i] - target))


def _build_grid_boundaries(cells: List[Dict[str, object]]) -> tuple[List[float], List[float]]:
    widths = [max(1.0, float(c["x2"]) - float(c["x1"])) for c in cells]
    heights = [max(1.0, float(c["y2"]) - float(c["y1"])) for c in cells]
    x_tol = max(2.0, float(median(widths)) * 0.25)
    y_tol = max(2.0, float(median(heights)) * 0.25)

    x_edges: List[float] = []
    y_edges: List[float] = []
    for c in cells:
        x_edges.extend([float(c["x1"]), float(c["x2"])])
        y_edges.extend([float(c["y1"]), float(c["y2"])])

    x_boundaries = _group_close_values(x_edges, x_tol)
    y_boundaries = _group_close_values(y_edges, y_tol)
    return x_boundaries, y_boundaries


def _group_rows(cells: List[Dict[str, object]]) -> List[List[Dict[str, object]]]:
    if not cells:
        return []
    heights = [float(c["y2"]) - float(c["y1"]) for c in cells]
    row_tol = max(3.0, float(median(heights)) * 0.6)
    ordered = sorted(cells, key=lambda c: (float(c["y1"]), float(c["x1"])))

    rows: List[List[Dict[str, object]]] = []
    current_row: List[Dict[str, object]] = []
    baseline_y = None
    for cell in ordered:
        y = float(cell["y1"])
        if baseline_y is None:
            current_row = [cell]
            baseline_y = y
            continue
        if abs(y - baseline_y) <= row_tol:
            current_row.append(cell)
            baseline_y = (baseline_y + y) / 2.0
        else:
            rows.append(sorted(current_row, key=lambda c: float(c["x1"])))
            current_row = [cell]
            baseline_y = y
    if current_row:
        rows.append(sorted(current_row, key=lambda c: float(c["x1"])))
    return rows


def export_to_excel(cells_with_text: List[Dict[str, object]], output_xlsx: str) -> str:
    if not cells_with_text:
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        wb.save(output_xlsx)
        print(f"[XLSX] 已输出: {output_xlsx}")
        return output_xlsx

    rows = _group_rows(cells_with_text)
    x_boundaries, y_boundaries = _build_grid_boundaries(cells_with_text)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    center_alignment = Alignment(horizontal="center", vertical="center")

    occupied = set()
    merge_count = 0
    skipped_overlap = 0

    # 优先放置跨行/跨列较大的单元格，避免与后续小框冲突。
    ordered_cells = sorted(
        cells_with_text,
        key=lambda c: (
            -((float(c["x2"]) - float(c["x1"])) * (float(c["y2"]) - float(c["y1"]))),
            float(c["y1"]),
            float(c["x1"]),
        ),
    )

    for cell in ordered_cells:
        x1, y1 = float(cell["x1"]), float(cell["y1"])
        x2, y2 = float(cell["x2"]), float(cell["y2"])
        c_start = _nearest_index(x_boundaries, x1)
        c_end = _nearest_index(x_boundaries, x2)
        r_start = _nearest_index(y_boundaries, y1)
        r_end = _nearest_index(y_boundaries, y2)

        if c_end <= c_start:
            c_end = min(c_start + 1, len(x_boundaries) - 1)
        if r_end <= r_start:
            r_end = min(r_start + 1, len(y_boundaries) - 1)

        if c_end <= c_start or r_end <= r_start:
            continue

        # 边界索引映射为 Excel 单元格索引
        col1, col2 = c_start + 1, c_end
        row1, row2 = r_start + 1, r_end

        overlap = False
        for rr in range(row1, row2 + 1):
            for cc in range(col1, col2 + 1):
                if (rr, cc) in occupied:
                    overlap = True
                    break
            if overlap:
                break
        if overlap:
            skipped_overlap += 1
            continue

        target_cell = ws.cell(row=row1, column=col1, value=str(cell.get("text", "")))
        target_cell.alignment = center_alignment
        if row2 > row1 or col2 > col1:
            ws.merge_cells(start_row=row1, start_column=col1, end_row=row2, end_column=col2)
            merge_count += 1
        for rr in range(row1, row2 + 1):
            for cc in range(col1, col2 + 1):
                occupied.add((rr, cc))

    wb.save(output_xlsx)
    print(f"[XLSX] merge_count={merge_count}, skipped_overlap={skipped_overlap}")
    print(f"[XLSX] 已输出: {output_xlsx}")
    return output_xlsx
